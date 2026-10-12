"""Valida que el Excel de la evaluación (fórmulas) dé lo mismo que el motor en Python.
Uso: python validar_excel.py cotizacion.json   (arma la evaluación, recalcula el libro con la librería formulas y compara)"""
from __future__ import annotations

import json
import sys
from pathlib import Path

import formulas
import openpyxl

AQUI = Path(__file__).parent
sys.path.insert(0, str(AQUI))
import evaluar  # noqa: E402


def main(path_cot: str) -> int:
    cot = json.loads(Path(path_cot).read_text(encoding="utf-8"))
    r = evaluar.correr(cot, con_radar=True, excel=True)
    xl = next(Path(r["carpeta"]).glob("Evaluacion_*.xlsx"))
    sol = formulas.ExcelModel().loads(str(xl)).finish().calculate()
    idx = {}
    for k in sol:
        ks = str(k).upper()
        if "]" in ks and "!" in ks:
            hoja, celda = ks.split("]", 1)[1].split("!", 1)
            idx[(hoja.strip("'"), celda)] = k

    def val(hoja, celda):
        k = idx.get((hoja.upper(), celda.upper().replace("$", "")))
        if k is None:
            return None
        v = sol[k].value[0, 0] if hasattr(sol[k], "value") else sol[k]
        return None if str(v) in ("", "empty") else v

    wb = openpyxl.load_workbook(xl)
    nombres = {n: d.attr_text for n, d in wb.defined_names.items()}

    def nombre(n):
        hoja, celda = nombres[n].split("!")
        return val(hoja.strip("'"), celda)

    difs, n = [], 0

    def cmp(etq, excel, py, tol=0.5):
        nonlocal n
        n += 1
        if isinstance(py, str) or isinstance(excel, str):
            ok = str(excel) == str(py)
        elif py is None or excel is None:
            ok = py is None and excel is None
        else:
            ok = abs(float(excel) - float(py)) <= max(tol, abs(float(py)) * 1e-9)
        if not ok:
            difs.append((etq, excel, py))

    imp, ch = r["importacion"], r["canales"]
    for nm, k in (("INTERNADO", "internado_clp"), ("COSTO_U", "costo_u"), ("CIF", "cif_usd"), ("IVA_IMP", "iva_import_clp")):
        cmp(nm, nombre(nm), imp[k], 0.01)
    for nm, k in (("PRECIO_PROM", "precio_prom"), ("CANAL_U", "canal_u")):
        cmp(nm, nombre(nm), ch[k], 0.01)
    cmp("OPER_U", nombre("OPER_U"), ch["manip_u"] + ch["insumos_u"] + ch["postventa_u"], 0.01)
    ws = wb["ROI"]
    filas = {ws.cell(i, 1).value: i for i in range(2, ws.max_row + 1)}
    mapa = {"Venta anual": "venta", "Costo de venta": "costo", "CONTRIBUCIÓN": "contribucion", "Contribución % de la venta": "contrib_pct",
            "CAPITAL EMPLEADO": "capital", "EVA (contribución − costo de capital)": "eva", "Capital: anticipo en producción": "k_anticipo",
            "Capital: pagado en tránsito + aduana": "k_transito", "Capital: inventario promedio": "k_bodega", "Capital: cuentas por cobrar": "k_cxc",
            "Crédito de flete e internación": "credito", "Almacenaje": "almacenaje", "ROI sobre capital": "roi"}
    for col, esc in (("B", "conservador"), ("C", "base"), ("D", "optimista")):
        for etq, k in mapa.items():
            cmp(f"ROI {esc} · {etq}", val("ROI", f"{col}{filas[etq]}"), r["anual"][esc][k], 1e-6 if k in ("contrib_pct", "roi") else 0.5)
    ws = wb["Sensibilidad"]
    for i, s in enumerate(r["sensibilidad"], 2):
        cmp(f"Sensibilidad {s['caso']} · costo", val("Sensibilidad", f"E{i}"), s["costo_u"], 0.01)
        cmp(f"Sensibilidad {s['caso']} · contrib/u", val("Sensibilidad", f"G{i}"), s["contrib_u"], 0.01)
    ws = wb["Flujo"]
    bloques = [i for i in range(1, ws.max_row + 1) if str(ws.cell(i, 1).value or "").startswith("Flujo del lote")]
    for esc, b0 in zip(("conservador", "base", "optimista"), bloques):
        met = {ws.cell(i, 1).value: i for i in range(b0, b0 + 80) if ws.cell(i, 1).value}
        lo = r["lote"][esc]
        for etq, k, tol in (("Utilidad del lote", "utilidad", 1), ("ROI del lote", "roi_lote", 1e-9), ("ROI anualizado", "roi_anualizado", 1e-9),
                            ("Caja máxima necesaria", "caja_necesaria", 1), ("Ciclo (meses, desde el anticipo)", "ciclo_meses", 0),
                            ("Unidades sin vender al mes máximo", "stock_sin_vender", 1e-6), ("GMROI del lote", "gmroi_lote", 1e-9)):
            cmp(f"Flujo {esc} · {etq}", val("Flujo", f"B{met[etq]}"), lo[k], tol)
        pb = val("Flujo", f"B{met['Payback (mes)']}")
        cmp(f"Flujo {esc} · payback", pb, lo["payback_mes"] if lo["payback_mes"] is not None else "no recupera", 0)
    cmp("Veredicto", val("Resumen", "B2"), r["veredicto_a_estos_precios"])
    print(f"{xl.name}: {n} comparaciones · diferencias {len(difs)}")
    for d in difs[:20]:
        print("  DIF", d)
    return 1 if difs else 0


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    sys.exit(main(sys.argv[1]))
